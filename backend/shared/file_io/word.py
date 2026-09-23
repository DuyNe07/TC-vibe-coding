"""Word (.docx) reading, writing and template rendering.

content = WordReader().read(uploaded_file)           # WordContent(paragraphs, tables)
doc = WordContent().add_heading("Báo cáo").add_paragraph("...").add_table(headers, rows)
data = WordWriter().write(doc)                        # bytes of a .docx
data = WordTemplateRenderer().render(template_path, {"name": "A"})   # docxtpl / Jinja2 {{ name }}
"""

from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from docx import Document
from docxtpl import DocxTemplate

from backend.core.exceptions import FileProcessingError
from backend.shared.file_io.base_file import BaseFileReader, BaseFileWriter, FileSource, read_source


@dataclass
class WordContent:
    """In-memory representation of a simple document (ordered blocks)."""

    paragraphs: list[str] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)
    blocks: list[tuple[str, Any]] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(self.paragraphs)

    def add_heading(self, text: str, level: int = 1) -> "WordContent":
        self.blocks.append(("heading", (text, level)))
        return self

    def add_paragraph(self, text: str) -> "WordContent":
        self.blocks.append(("paragraph", text))
        return self

    def add_table(self, headers: list[str], rows: list[list[Any]]) -> "WordContent":
        self.blocks.append(("table", (headers, rows)))
        return self


class WordReader(BaseFileReader[WordContent]):
    supported_extensions = (".docx",)

    def _read(self, stream: BytesIO, extension: str) -> WordContent:
        document = Document(stream)
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        tables = [[[cell.text.strip() for cell in row.cells] for row in table.rows] for table in document.tables]
        return WordContent(paragraphs=paragraphs, tables=tables)


class WordWriter(BaseFileWriter[WordContent]):
    def write(self, data: WordContent) -> bytes:
        document = Document()
        for kind, value in data.blocks:
            if kind == "heading":
                document.add_heading(value[0], level=value[1])
            elif kind == "paragraph":
                document.add_paragraph(value)
            elif kind == "table":
                headers, rows = value
                table = document.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                for cell, header in zip(table.rows[0].cells, headers, strict=True):
                    cell.text = str(header)
                for row in rows:
                    for cell, item in zip(table.add_row().cells, row, strict=False):
                        cell.text = "" if item is None else str(item)
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()


class WordTemplateRenderer:
    """Fills a .docx template containing Jinja2 placeholders (``{{ field }}``, ``{% for %}``)."""

    def render(self, template: FileSource, context: dict[str, Any]) -> bytes:
        content, _ = read_source(template)
        try:
            document = DocxTemplate(BytesIO(content))
            document.render(context)
            buffer = BytesIO()
            document.save(buffer)
        except Exception as exc:
            raise FileProcessingError(f"Cannot build the Word file from the template: {exc}") from exc
        return buffer.getvalue()
