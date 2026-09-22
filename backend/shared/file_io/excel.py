"""Excel / CSV reading and writing.

rows = ExcelReader().read(uploaded_file)              # list[dict] (header row -> keys)
content = ExcelWriter().write({"Sheet1": rows})       # bytes of a styled .xlsx
"""

from collections.abc import Mapping, Sequence
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.shared.file_io.base_file import BaseFileReader, BaseFileWriter, FileSource, read_source

Rows = list[dict[str, Any]]
SheetData = Sequence[Mapping[str, Any]] | pd.DataFrame


class ExcelReader(BaseFileReader[Rows]):
    """Reads one sheet into a list of dicts. Empty cells become ``None``; fully empty rows are dropped."""

    supported_extensions = (".xlsx", ".xlsm", ".xls", ".csv")

    def __init__(self, sheet_name: str | int = 0, header_row: int = 0) -> None:
        self.sheet_name = sheet_name
        self.header_row = header_row

    def read_dataframe(self, source: FileSource) -> pd.DataFrame:
        return pd.DataFrame(self.read(source))

    def sheet_names(self, source: FileSource) -> list[str]:
        content, _ = read_source(source)
        return list(pd.ExcelFile(BytesIO(content)).sheet_names)

    def _read(self, stream: BytesIO, extension: str) -> Rows:
        if extension == ".csv":
            frame = pd.read_csv(stream, header=self.header_row, dtype=object, encoding="utf-8-sig")
        else:
            engine = "xlrd" if extension == ".xls" else "openpyxl"
            frame = pd.read_excel(
                stream, sheet_name=self.sheet_name, header=self.header_row, dtype=object, engine=engine
            )
        frame = frame.dropna(how="all")
        frame.columns = [str(col).strip() for col in frame.columns]
        frame = frame.astype(object).where(pd.notna(frame), None)
        return frame.to_dict(orient="records")


class ExcelWriter(BaseFileWriter[Mapping[str, SheetData]]):
    """Writes ``{sheet_name: rows_or_dataframe}`` to a styled .xlsx (bold header, frozen, auto width)."""

    header_fill = PatternFill("solid", fgColor="0F62FE")
    header_font = Font(bold=True, color="FFFFFF")

    def write(self, data: Mapping[str, SheetData]) -> bytes:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name, sheet_data in data.items():
            rows = sheet_data.to_dict(orient="records") if isinstance(sheet_data, pd.DataFrame) else list(sheet_data)
            self._write_sheet(workbook.create_sheet(title=sheet_name[:31]), rows)
        if not workbook.sheetnames:
            workbook.create_sheet("Sheet1")
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _write_sheet(self, sheet: Any, rows: list[Mapping[str, Any]]) -> None:
        headers = list(rows[0].keys()) if rows else []
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(h) for h in headers])
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index)
            cell.fill, cell.font = self.header_fill, self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            width = max([len(str(header))] + [len(str(r.get(header) or "")) for r in rows[:500]])
            sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 60)
        sheet.freeze_panes = "A2"
