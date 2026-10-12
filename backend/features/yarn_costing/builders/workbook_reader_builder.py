"""UC-34 uploaded workbook → rows of the 7 sheets of 5.2 (header found in the first 10 rows, values of the cells).

openpyxl / xlrd are used directly: the shared ExcelReader reads one fixed header row only, while this file format
needs a header spread over two rows, the second "Mã Kế toán" column and (S1) the formula flag of loss cells.
"""

import re
from collections.abc import Callable, Iterator
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

import openpyxl
import xlrd

from backend.core.base import BaseBuilder, UploadedFileDTO
from backend.features.yarn_costing.constants import (
    BLEND_SLOTS,
    COL_ACCOUNTING,
    COL_COMPONENT_COUNT,
    COL_MAN_DAYS,
    COL_MATERIAL,
    COL_MINUTES,
    COL_NOTE,
    COL_PACKAGING_CODE,
    COL_PACKAGING_NAME,
    COL_PRICED_DATE,
    COL_REMARK,
    COL_ROUTE,
    COL_ROUTE_TECH,
    COL_STAGE,
    COL_STAGE_GROUP,
    COL_STAGE_NAME,
    COL_TECHNOLOGY,
    COL_TOTAL,
    COL_WEIGHT,
    COL_YARN_CODE,
    COL_YARN_GROUP,
    COL_YARN_NAME,
    COST_COLUMNS,
    EMPTY_ROWS_STOP,
    HEADER_SCAN_ROWS,
    RATIO_SUFFIX,
    RECOVERY_COLUMN,
    SHEET_LOSSES,
    SHEET_NSM,
    SHEET_ORDER,
    SHEET_PACKAGING,
    SHEET_PRICES,
    SHEET_RATES,
    SHEET_RECOVERY,
    SHEET_ROUTES,
)
from backend.features.yarn_costing.exceptions import ImportFileError
from backend.features.yarn_costing.models.enums import FactoryCode
from backend.features.yarn_costing.models.import_rows import (
    HeaderMap,
    LossRow,
    NsmRow,
    PackagingRow,
    ParsedWorkbook,
    PriceRow,
    RateRow,
    RecoveryRow,
    RouteRow,
    SheetRead,
)
from backend.features.yarn_costing.models.yarn import CostBreakdown
from backend.shared.utils import normalize_key, parse_decimal

RowSource = Callable[[int, int | None], Iterator[tuple[Any, ...]]]
LABEL_COLUMN = 1  # "Hoi phe": row labels are in column B
BAD = object()


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def number(value: Any) -> Any:
    """Decimal, None for empty cells, BAD for text that is not a number (e.g. '#REF!')."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    parsed = parse_decimal(value)
    return BAD if parsed is None else parsed


def to_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = re.fullmatch(r"\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})\s*", str(value or ""))
    if not match:
        return None
    day, month, year = (int(g) for g in match.groups())
    try:
        return date(year, month, day)
    except ValueError:
        return None


def find_header(rows: list[tuple[Any, ...]], names: list[str]) -> HeaderMap:
    """5.1: search the wanted column names / stage codes in the first rows; a column may be named on two rows."""
    wanted = {normalize_key(n) for n in names}
    needed = min(2, len(wanted))  # a header row names at least 2 different wanted columns (data rows rarely do)
    columns: dict[str, list[int]] = {}
    last_row = 0
    for row_no, row in enumerate(rows, start=1):
        hits = [
            (col, key)
            for col, value in enumerate(row)
            if isinstance(value, str) and (key := normalize_key(value)) in wanted
        ]
        if len({key for _, key in hits}) < needed:
            continue
        for col, key in hits:
            if col not in columns.setdefault(key, []):
                columns[key].append(col)
        last_row = max(last_row, row_no)
    return HeaderMap(columns={k: sorted(v) for k, v in columns.items()}, last_row=last_row)


def col_of(header: HeaderMap, name: str, occurrence: int = 0) -> int | None:
    return header.col(normalize_key(name), occurrence)


def missing_of(header: HeaderMap, names: list[str]) -> list[str]:
    return [n for n in names if col_of(header, n) is None]


class WorkbookReaderBuilder(BaseBuilder[ParsedWorkbook]):
    def __init__(self, file: UploadedFileDTO, factory: FactoryCode, stage_codes: list[str]) -> None:
        self.file, self.factory, self.stage_codes = file, factory, stage_codes
        self.slots = list(BLEND_SLOTS[factory.value])
        self.sheets: list[SheetRead] = []

    # ---- workbook access ---------------------------------------------------------------------------------------
    def _open(self, formulas: bool = False) -> tuple[dict[str, RowSource], Callable[[], None]]:
        canonical = {normalize_key(name): name for name in SHEET_ORDER}
        sources: dict[str, RowSource] = {}
        try:
            if self.file.extension == ".xls":
                book = xlrd.open_workbook(file_contents=self.file.content)
                for sheet in book.sheets():
                    if (name := canonical.get(normalize_key(sheet.name))) is not None:
                        sources[name] = self._xls_source(book, sheet)
                return sources, lambda: None
            workbook = openpyxl.load_workbook(BytesIO(self.file.content), read_only=True, data_only=not formulas)
        except Exception as exc:
            raise ImportFileError(f"Không đọc được file '{self.file.filename}'. Hãy kiểm tra lại file Excel.") from exc
        exact = {" ".join(name.split()).casefold(): name for name in SHEET_ORDER}
        titles = sorted(workbook.sheetnames, key=lambda t: " ".join(t.split()).casefold() not in exact)
        for title in titles:  # exact names first ('Hoi phe' wins over another sheet named 'Hồi Phế')
            if (name := canonical.get(normalize_key(title))) is not None and name not in sources:
                ws = workbook[title]
                sources[name] = lambda start, max_col, ws=ws: ws.iter_rows(
                    min_row=start, max_col=max_col, values_only=True
                )
        return sources, workbook.close

    def _xls_source(self, book: Any, sheet: Any) -> RowSource:
        def rows(start: int, max_col: int | None) -> Iterator[tuple[Any, ...]]:
            for r in range(start - 1, sheet.nrows):
                values = []
                for c in range(min(sheet.ncols, max_col or sheet.ncols)):
                    cell = sheet.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        values.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                    elif cell.ctype == xlrd.XL_CELL_EMPTY:
                        values.append(None)
                    else:
                        values.append(cell.value)
                yield tuple(values)

        return rows

    def _data(
        self, source: RowSource, header: HeaderMap, key_cols: list[int | None]
    ) -> Iterator[tuple[int, tuple[Any, ...]]]:
        """Rows after the last header row; stops after many empty rows; skips column-numbering rows."""
        empty = 0
        keys = [c for c in key_cols if c is not None]
        width = header.max_col
        for row_no, row in enumerate(source(header.last_row + 1, width), start=header.last_row + 1):
            row = tuple(row) + (None,) * max(0, width - len(row))
            values = [row[c] for c in keys]
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                empty += 1
                if empty >= EMPTY_ROWS_STOP:
                    break
                continue
            empty = 0
            if all(isinstance(v, int | float) and not isinstance(v, bool) for v in values if v is not None):
                continue  # a row of column numbers (1, 2, 3...) under the header
            yield row_no, row

    def _header(self, source: RowSource, names: list[str]) -> HeaderMap:
        rows = [tuple(r) for _, r in zip(range(HEADER_SCAN_ROWS), source(1, None), strict=False)]
        return find_header(rows, names)

    def _register(self, sheet: str, header: HeaderMap, required: list[str], count: int) -> None:
        self.sheets.append(SheetRead(sheet=sheet, read=count, missing_columns=missing_of(header, required)))

    # ---- build -------------------------------------------------------------------------------------------------
    def build(self) -> ParsedWorkbook:
        sources, close = self._open()
        try:
            result: dict[str, list] = {}
            readers = {
                SHEET_RATES: self._rates,
                SHEET_RECOVERY: self._recovery,
                SHEET_PACKAGING: self._packaging,
                SHEET_ROUTES: self._routes,
                SHEET_LOSSES: self._losses,
                SHEET_NSM: self._nsm,
                SHEET_PRICES: self._prices,
            }
            for sheet in SHEET_ORDER:
                if sheet in sources:
                    result[sheet] = readers[sheet](sources[sheet])
        finally:
            close()
        if SHEET_LOSSES in sources and self.factory == FactoryCode.S1 and self.file.extension != ".xls":
            result[SHEET_LOSSES] = self._mark_formulas(result[SHEET_LOSSES])
        return ParsedWorkbook(
            file_name=self.file.filename,
            sheets=sorted(self.sheets, key=lambda s: SHEET_ORDER.index(s.sheet)),
            rate_rows=result.get(SHEET_RATES, []),
            recovery_rows=result.get(SHEET_RECOVERY, []),
            packaging_rows=result.get(SHEET_PACKAGING, []),
            route_rows=result.get(SHEET_ROUTES, []),
            loss_rows=result.get(SHEET_LOSSES, []),
            nsm_rows=result.get(SHEET_NSM, []),
            price_rows=result.get(SHEET_PRICES, []),
        )

    def _cost(self, row: tuple[Any, ...], cols: list[int | None]) -> list[Any]:
        return [number(row[c]) if c is not None else None for c in cols]

    def _rates(self, source: RowSource) -> list[RateRow]:
        names = [COL_STAGE, COL_STAGE_NAME, COL_STAGE_GROUP, *COST_COLUMNS, COL_NOTE]
        header = self._header(source, names)
        stage, group, note = col_of(header, COL_STAGE), col_of(header, COL_STAGE_GROUP), col_of(header, COL_NOTE)
        costs = [col_of(header, n) for n in COST_COLUMNS]
        rows: list[RateRow] = []
        if stage is not None:
            for row_no, row in self._data(source, header, [stage]):
                values = [None if v is BAD else v for v in self._cost(row, costs)]
                bad = any(v is BAD for v in self._cost(row, costs))
                rows.append(
                    RateRow(
                        excel_row=row_no,
                        stage_value=text(row[stage]),
                        group_value=text(row[group]) if group is not None else "",
                        labor=None if bad else values[0],
                        depreciation=values[1],
                        electricity=values[2],
                        materials=values[3],
                        overhead=values[4],
                        note=text(row[note]) if note is not None else "",
                    )
                )
        self._register(SHEET_RATES, header, [COL_STAGE, *COST_COLUMNS], len(rows))
        return rows

    def _recovery(self, source: RowSource) -> list[RecoveryRow]:
        column = RECOVERY_COLUMN[self.factory.value]
        header = self._header(source, [column])
        col = col_of(header, column)
        rows: list[RecoveryRow] = []
        if col is not None:
            width = max(col, LABEL_COLUMN) + 1
            for row_no, row in enumerate(source(header.last_row + 1, width), start=header.last_row + 1):
                row = tuple(row) + (None,) * width
                label, value = text(row[LABEL_COLUMN]), number(row[col])
                if label:
                    rows.append(RecoveryRow(excel_row=row_no, label=label, rate=None if value is BAD else value))
                if row_no > header.last_row + 50:
                    break
        self._register(SHEET_RECOVERY, header, [column], len(rows))
        return rows

    def _packaging(self, source: RowSource) -> list[PackagingRow]:
        names = [COL_PACKAGING_NAME, COL_PACKAGING_CODE, COL_WEIGHT, COL_MAN_DAYS, COL_MINUTES]
        header = self._header(source, names)
        code = col_of(header, COL_PACKAGING_CODE)
        cols = {n: col_of(header, n) for n in names}
        rows: list[PackagingRow] = []
        if code is not None:
            for row_no, row in self._data(source, header, [code]):
                nums = {n: number(row[c]) if c is not None else None for n, c in cols.items()}
                nums = {n: None if v is BAD else v for n, v in nums.items()}
                name_col = cols[COL_PACKAGING_NAME]
                rows.append(
                    PackagingRow(
                        excel_row=row_no,
                        code=text(row[code]),
                        name=text(row[name_col]) if name_col is not None else "",
                        weight_kg=nums[COL_WEIGHT],
                        man_days=nums[COL_MAN_DAYS],
                        minutes_per_kg=nums[COL_MINUTES],
                    )
                )
        self._register(SHEET_PACKAGING, header, [COL_PACKAGING_CODE, COL_MINUTES], len(rows))
        return rows

    def _routes(self, source: RowSource) -> list[RouteRow]:
        names = [COL_YARN_GROUP, COL_ROUTE_TECH, COL_COMPONENT_COUNT, COL_ROUTE]
        header = self._header(source, names)
        group, tech, count, route = (col_of(header, n) for n in names)
        rows: list[RouteRow] = []
        if group is not None:
            for row_no, row in self._data(source, header, [group]):
                value = number(row[count]) if count is not None else None
                rows.append(
                    RouteRow(
                        excel_row=row_no,
                        yarn_group=text(row[group]),
                        technology_text=text(row[tech]) if tech is not None else "",
                        component_count=None if value is BAD else value,
                        route_text=text(row[route]) if route is not None else "",
                    )
                )
        self._register(SHEET_ROUTES, header, [COL_YARN_GROUP, COL_ROUTE_TECH], len(rows))
        return rows

    def _losses(self, source: RowSource) -> list[LossRow]:
        with_components = self.factory == FactoryCode.S4
        slot_names = self.slots if with_components else []
        ratio_names = [s + RATIO_SUFFIX for s in slot_names]
        names = [COL_YARN_GROUP, COL_MATERIAL, COL_TECHNOLOGY, *self.stage_codes, *slot_names, *ratio_names]
        header = self._header(source, names)
        group, material, tech = (
            col_of(header, COL_YARN_GROUP),
            col_of(header, COL_MATERIAL),
            col_of(header, COL_TECHNOLOGY),
        )
        stage_cols = {code: col_of(header, code) for code in self.stage_codes}
        rows: list[LossRow] = []
        if group is not None:
            for row_no, row in self._data(source, header, [group]):
                losses, bad = {}, []
                for code, col in stage_cols.items():
                    if col is None:
                        continue
                    value = number(row[col])
                    if value is BAD:
                        bad.append(code)
                    elif value is not None:
                        losses[code] = value
                comp_names = [text(row[c]) if (c := col_of(header, n)) is not None else "" for n in slot_names]
                comp_ratios = [number(row[c]) if (c := col_of(header, n)) is not None else None for n in ratio_names]
                rows.append(
                    LossRow(
                        excel_row=row_no,
                        yarn_group_label=text(row[group]),
                        material=text(row[material]) if material is not None else "",
                        technology_text=text(row[tech]) if tech is not None else "",
                        losses=losses,
                        component_names=comp_names,
                        component_ratios=[None if v is BAD else v for v in comp_ratios],
                        bad_values=bad,
                    )
                )
        self._register(SHEET_LOSSES, header, [COL_YARN_GROUP, COL_TECHNOLOGY], len(rows))
        self._loss_header = header
        return rows

    def _mark_formulas(self, rows: list[LossRow]) -> list[LossRow]:
        """BR-42 (S1): a loss cell holding a formula → group 'Nhập từ file - chưa khai báo thành phần'."""
        sources, close = self._open(formulas=True)
        try:
            header = self._loss_header
            cols = [c for code in self.stage_codes if (c := col_of(header, code)) is not None]
            formula_rows: set[int] = set()
            wanted = {r.excel_row for r in rows}
            last = max(wanted, default=0)
            for row_no, row in enumerate(
                sources[SHEET_LOSSES](header.last_row + 1, header.max_col), start=header.last_row + 1
            ):
                if row_no > last:
                    break
                row = tuple(row) + (None,) * header.max_col
                if row_no in wanted and any(isinstance(row[c], str) and row[c].startswith("=") for c in cols):
                    formula_rows.add(row_no)
        finally:
            close()
        return [r.model_copy(update={"has_formula": True}) if r.excel_row in formula_rows else r for r in rows]

    def _nsm(self, source: RowSource) -> list[NsmRow]:
        names = [COL_YARN_CODE, COL_YARN_NAME, *self.slots, *self.stage_codes, COL_PACKAGING_NAME, COL_REMARK]
        header = self._header(source, names)
        code, name, packaging = (
            col_of(header, COL_YARN_CODE),
            col_of(header, COL_YARN_NAME),
            col_of(header, COL_PACKAGING_NAME),
        )
        remark = col_of(header, COL_REMARK)
        slot_cols = {s: col_of(header, s) for s in self.slots}
        stage_cols = {s: col_of(header, s) for s in self.stage_codes}
        rows: list[NsmRow] = []
        if name is not None:
            for row_no, row in self._data(source, header, [name]):
                bad: list[str] = []
                blend: dict[str, Decimal] = {}
                nsm: dict[str, Decimal] = {}
                for slot, col in slot_cols.items():
                    value = number(row[col]) if col is not None else None
                    if value is BAD:
                        bad.append(slot)
                    elif value is not None:
                        blend[slot] = value
                for stage, col in stage_cols.items():
                    value = number(row[col]) if col is not None else None
                    if value is BAD or (value is not None and value < 0):
                        bad.append(stage)
                    elif value is not None:
                        nsm[stage] = value
                rows.append(
                    NsmRow(
                        excel_row=row_no,
                        raw_code=text(row[code]) if code is not None else "",
                        yarn_name=text(row[name]),
                        blend_ratios=blend,
                        nsm=nsm,
                        packaging_code=text(row[packaging]) if packaging is not None else "",
                        remark=text(row[remark]) if remark is not None else "",
                        bad_values=bad,
                    )
                )
        self._register(SHEET_NSM, header, [COL_YARN_NAME], len(rows))
        return rows

    def _prices(self, source: RowSource) -> list[PriceRow]:
        names = [COL_ACCOUNTING, *COST_COLUMNS, COL_TOTAL, COL_PRICED_DATE]
        header = self._header(source, names)
        code, accounting = col_of(header, COL_ACCOUNTING, 0), col_of(header, COL_ACCOUNTING, 1)
        costs = [col_of(header, n) for n in COST_COLUMNS]
        total_col, date_col = col_of(header, COL_TOTAL), col_of(header, COL_PRICED_DATE)
        rows: list[PriceRow] = []
        if code is not None:
            for row_no, row in self._data(source, header, [code]):
                values = [v if v is not BAD and v is not None else Decimal(0) for v in self._cost(row, costs)]
                total = number(row[total_col]) if total_col is not None else None
                raw_date = row[date_col] if date_col is not None else None
                rows.append(
                    PriceRow(
                        excel_row=row_no,
                        raw_code=text(row[code]),
                        cost=CostBreakdown(
                            labor=values[0],
                            depreciation=values[1],
                            electricity=values[2],
                            materials=values[3],
                            overhead=values[4],
                        ),
                        total=None if total is BAD else total,
                        priced_date=to_date(raw_date),
                        priced_date_text=raw_date.strftime("%d/%m/%Y")
                        if isinstance(raw_date, datetime | date)
                        else text(raw_date),
                        accounting_code=text(row[accounting]) if accounting is not None else "",
                    )
                )
        self._register(SHEET_PRICES, header, [COL_ACCOUNTING], len(rows))
        return rows
