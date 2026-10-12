"""5.7 File khai báo thành phần sợi pha: read the uploaded sheet (UC-45) and write the template (UC-44).

openpyxl / xlrd are used directly (like workbook_reader_builder): the sheet is found by the start of its name and the
template needs a drop-down list of the groups.
"""

from io import BytesIO
from typing import Any

import openpyxl
import xlrd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from backend.core.base import BaseBuilder, FileDownloadDTO, MimeType, UploadedFileDTO
from backend.features.yarn_costing.constants import (
    COL_COMP_GROUP,
    COL_COMP_LEVEL,
    COL_COMP_LOSSES,
    COL_COMP_NAME,
    COL_COMP_NO,
    COL_COMP_RATIO,
    COL_COMP_TECH,
    COL_COMP_TOTAL,
    COL_COMP_USAGE,
    COMP_GUIDE_LINES,
    COMP_GUIDE_SHEET,
    COMP_MAX,
    COMP_REF_SHEET,
    COMP_SHEET,
    MSG_COMP_MISSING_COLUMNS,
    MSG_COMP_NO_SHEET,
)
from backend.features.yarn_costing.exceptions import ImportFileError
from backend.features.yarn_costing.models.enums import FactoryCode
from backend.features.yarn_costing.models.import_rows import ComponentDeclRow
from backend.features.yarn_costing.models.loss_group import LossGroup
from backend.shared.utils import normalize_key, parse_decimal

NAME_COLUMNS = [COL_COMP_NAME.format(k=k) for k in range(1, COMP_MAX + 1)]
RATIO_COLUMNS = [COL_COMP_RATIO.format(k=k) for k in range(1, COMP_MAX + 1)]
REQUIRED_COLUMNS = [COL_COMP_GROUP, NAME_COLUMNS[0], RATIO_COLUMNS[0]]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _percent(value: Any) -> tuple[Any, bool]:
    """Ratio cell → (Decimal percent | None, is_bad). '48', 48, '48%' and '48,5' are read; text that is not a number is bad."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, False
    if isinstance(value, str):
        value = value.strip().rstrip("%").strip().replace(",", ".")
    number = parse_decimal(value)
    return (number, False) if number is not None else (None, True)


def _is_component_sheet(title: str) -> bool:
    return normalize_key(title).startswith(normalize_key(COMP_SHEET))


class ComponentFileReaderBuilder(BaseBuilder[list[ComponentDeclRow]]):
    """UC-45: sheet 'Khai bao thanh phan…' → rows (header in row 1, columns by name; other columns ignored)."""

    def __init__(self, file: UploadedFileDTO) -> None:
        self.file = file

    def _rows(self) -> list[tuple[Any, ...]]:
        try:
            if self.file.extension == ".xls":
                book = xlrd.open_workbook(file_contents=self.file.content)
                sheet = next((s for s in book.sheets() if _is_component_sheet(s.name)), None)
                if sheet is None:
                    raise ImportFileError(MSG_COMP_NO_SHEET)
                return [tuple(sheet.row_values(i)) for i in range(sheet.nrows)]
            workbook = openpyxl.load_workbook(BytesIO(self.file.content), read_only=True, data_only=True)
        except ImportFileError:
            raise
        except Exception as exc:
            raise ImportFileError(f"Không đọc được file '{self.file.filename}'. Hãy kiểm tra lại file Excel.") from exc
        try:
            title = next((t for t in workbook.sheetnames if _is_component_sheet(t)), None)
            if title is None:
                raise ImportFileError(MSG_COMP_NO_SHEET)
            return [tuple(r) for r in workbook[title].iter_rows(values_only=True)]
        finally:
            workbook.close()

    def build(self) -> list[ComponentDeclRow]:
        rows = self._rows()
        header = [normalize_key(_text(c)) for c in rows[0]] if rows else []
        index = {
            name: header.index(normalize_key(name))
            for name in [COL_COMP_GROUP, COL_COMP_TECH, *NAME_COLUMNS, *RATIO_COLUMNS]
            if normalize_key(name) in header
        }
        missing = [name for name in REQUIRED_COLUMNS if name not in index]
        if missing:
            raise ImportFileError(MSG_COMP_MISSING_COLUMNS.format(columns=", ".join(missing)))

        def cell(row: tuple[Any, ...], name: str) -> Any:
            col = index.get(name)
            return row[col] if col is not None and col < len(row) else None

        result: list[ComponentDeclRow] = []
        for row_no, row in enumerate(rows[1:], start=2):
            label = _text(cell(row, COL_COMP_GROUP))
            names = [_text(cell(row, name)) for name in NAME_COLUMNS]
            if not label and not any(names):
                continue
            ratios, bad = [], []
            for k, name in enumerate(RATIO_COLUMNS, start=1):
                value, is_bad = _percent(cell(row, name))
                ratios.append(value)
                if is_bad:
                    bad.append(k)
            result.append(
                ComponentDeclRow(
                    excel_row=row_no,
                    group_label=label,
                    technology_text=_text(cell(row, COL_COMP_TECH)),
                    names=names,
                    ratios=ratios,
                    bad_ratios=bad,
                )
            )
        return result


def _loss_text(group: LossGroup) -> str:
    return "; ".join(f"{code} {float(value * 100):g}" for code, value in group.losses.items() if value)


class ComponentTemplateBuilder(BaseBuilder[FileDownloadDTO]):
    """UC-44: mau-khai-bao-thanh-phan-<xưởng>.xlsx (5.7) - one row per loss group, secondary groups prefilled."""

    def __init__(self, factory: FactoryCode, groups: list[LossGroup], usage: dict[str, int]) -> None:
        self.factory, self.groups, self.usage = factory, groups, usage

    def build(self) -> FileDownloadDTO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = COMP_SHEET
        headers = [COL_COMP_NO, COL_COMP_GROUP, COL_COMP_TECH, COL_COMP_LEVEL, COL_COMP_USAGE]
        for name, ratio in zip(NAME_COLUMNS, RATIO_COLUMNS, strict=True):
            headers += [name, ratio]
        headers += [COL_COMP_TOTAL, COL_COMP_LOSSES]
        ws.append(headers)
        first_name_col = headers.index(NAME_COLUMNS[0]) + 1
        ordered = sorted(self.groups, key=lambda g: (-self.usage.get(g.group_key, 0), g.group_key))
        for number, group in enumerate(ordered, start=1):
            pairs: list[Any] = []
            for component in group.components[:COMP_MAX]:
                pairs += [component.component_group_key, float(component.ratio * 100)]
            pairs += [None] * (2 * COMP_MAX - len(pairs))
            ws.append(
                [
                    number,
                    group.yarn_group_label or group.group_key,
                    group.technology.value,
                    group.level.value,
                    self.usage.get(group.group_key, 0),
                    *pairs,
                    None,
                    _loss_text(group),
                ]
            )
            r = ws.max_row
            ratio_cells = ",".join(f"{get_column_letter(first_name_col + 2 * k + 1)}{r}" for k in range(COMP_MAX))
            ws.cell(row=r, column=first_name_col + 2 * COMP_MAX, value=f"=SUM({ratio_cells})")
        ref = wb.create_sheet(COMP_REF_SHEET)
        ref.append(["Nhóm tiêu hao", "Loại", "Nguyên liệu", "Công nghệ", "Tiêu hao (%)"])
        for group in sorted(self.groups, key=lambda g: g.group_key):
            ref.append([group.group_key, group.level.value, group.material, group.technology.value, _loss_text(group)])
        if self.groups:
            choices = DataValidation(
                type="list", formula1=f"='{COMP_REF_SHEET}'!$A$2:$A${ref.max_row}", allow_blank=True
            )
            choices.showErrorMessage = False
            ws.add_data_validation(choices)
            last_row = max(ws.max_row, 2) + 200  # room for new groups typed at the end
            for k in range(COMP_MAX):
                col = get_column_letter(first_name_col + 2 * k)
                choices.add(f"{col}2:{col}{last_row}")
        guide = wb.create_sheet(COMP_GUIDE_SHEET)
        for line in COMP_GUIDE_LINES:
            guide.append([line])
        guide.column_dimensions["A"].width = 140
        for sheet in (ws, ref):
            for c in sheet[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="DDEBF7")
                c.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.freeze_panes = "C2" if sheet is ws else "A2"
            for i, column in enumerate(sheet.columns, start=1):
                width = max(len(str(c.value or "")) for c in list(column)[:300]) + 2
                sheet.column_dimensions[get_column_letter(i)].width = min(45, max(9, width))
        buffer = BytesIO()
        wb.save(buffer)
        return FileDownloadDTO(
            filename=f"mau-khai-bao-thanh-phan-{self.factory.value}.xlsx",
            content=buffer.getvalue(),
            mime_type=MimeType.XLSX,
        )
