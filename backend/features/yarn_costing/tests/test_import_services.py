"""Service tests of UC-33, UC-34, UC-35, UC-41 with workbooks built in memory (same layout as the real files)."""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from backend.core.base import UploadedFileDTO
from backend.core.exceptions import AppError, BusinessRuleViolation
from backend.features.yarn_costing.dto.import_dto import (
    DownloadTemplateRequest,
    GetLegacyComparisonRequest,
    ImportWorkbookRequest,
    ListImportBatchesRequest,
)
from backend.features.yarn_costing.dto.pricing_dto import GetYarnDetailRequest
from backend.features.yarn_costing.models.enums import FactoryCode, ImportMode, LossGroupLevel, YarnStatus
from backend.features.yarn_costing.repositories.loss_group_repository import LossGroupRepository
from backend.features.yarn_costing.services.download_template_service import DownloadTemplateService
from backend.features.yarn_costing.services.get_legacy_comparison_service import GetLegacyComparisonService
from backend.features.yarn_costing.services.get_yarn_detail_service import GetYarnDetailService
from backend.features.yarn_costing.services.import_workbook_service import ImportWorkbookService
from backend.features.yarn_costing.services.list_import_batches_service import ListImportBatchesService
from backend.features.yarn_costing.tests.helpers import S4_FILE_NSM, S4_LOSSES, S4_NAME, S4_RATES, SO, add_member

D = Decimal
S4 = FactoryCode.S4
S4_STAGES = [
    "C1",
    "C2",
    "C3.1",
    "C3.2",
    "C3.3",
    "SBD",
    "PD",
    "PD1",
    "PCD1",
    "PCD2",
    "CM",
    "DCM",
    "MD1",
    "MD2",
    "MD3",
    "R",
    "S",
    "S Compact",
    "W",
    "V",
]
SLOTS = ["PE", "Other - C3.1", "Other - C3.2", "Other - C3.3", "Cotton"]
# The old file declares the C3 loss of S4-2017 in column C3.3 (BR-59 reads it as C3)
FILE_LOSSES = {("C3.3" if code == "C3" else code): value for code, value in S4_LOSSES.items()}


def s4_workbook(
    extra_nsm: list[list] | None = None, sheets: set[str] | None = None, duplicate_group: bool = False
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name: str):
        return wb.create_sheet(name) if sheets is None or name in sheets else Workbook().active

    ws = sheet("Dinh muc (phut)")
    ws.append(["Bảng định mức chi phí/ phút"])
    ws.append([])
    ws.append(
        [
            "Stt",
            "Xưởng sợi",
            "Công đoạn",
            "Tên công đoạn",
            "Nhóm công đoạn",
            "Nhân công",
            "Khấu hao",
            "Điện",
            "Vật tư",
            "CPQL",
            "Giá thành sản xuất",
        ]
    )
    ws.append(list(range(1, 12)))  # a row of column numbers
    for i, code in enumerate([*S4_STAGES, "C3", "Type 3"], 1):
        values = S4_RATES.get(code, S4_RATES["MD1"])
        ws.append([i, "S4", code, code, code, *(float(v) for v in values), None])
    ws = sheet("Hoi phe")
    for _ in range(4):
        ws.append([])
    ws.append([None, "Tên công đoạn", "Sợi 1 ", "Sợi 4 ", "Công đoạn"])
    ws.append(["x", "Cúi ghép sau máy thô ", 0.02, 0.015, "C"])
    ws.append(["x", "Bông A2 ", 0.01, 0.025, "S"])
    ws.append(["x", "Sợi thô ", 0.01, 0.01, "R"])
    ws = sheet("Dong goi")
    ws.append(["Stt", "Loại đóng gói", "Loại", "Khối lượng", "Công", "Phút/kg"])
    ws.append([1, "Xe tải 5T", "Type 3", 4989.6, 2.28125, 0.2194564694564695])
    ws = sheet("QTCN theo nhom soi")
    ws.append(["Stt", "Xưởng sợi", "Nhóm sợi", "Phân loại", "Số thành phần", "Quy trình công nghệ"])
    ws.append([1, "S4", "CORE", "MVS", 2, "B1 - C1 - C3 - CM - MD1 - MD2 - MD3 - V - F"])
    ws = sheet("2. Input tieu hao")
    ws.append(["Bảng tiêu hao"])
    ws.append([])
    ws.append(["Stt", "Nhóm sợi", "Nguyên liệu", "Công nghệ", "tloc", "Loại sợi"])
    ws.append(["Stt", "Nhóm sợi", "Nguyên liệu", "Công nghệ", "tloc", *SLOTS, *(s + "2" for s in SLOTS), *S4_STAGES])
    core = [1, "CORE (COTTON/ FDY 30D/24F/1) (70/30) MVS", "Core", "mvs", None] + [None] * 10
    core += [float(FILE_LOSSES[s]) if s in FILE_LOSSES else None for s in S4_STAGES]
    ws.append(core)
    pe = (
        [2, "PE", "PE Normal", "Ring", None]
        + [None] * 10
        + [0.025 if s == "C2" else 0.01 if s == "W" else None for s in S4_STAGES]
    )
    vi = [3, "VISCOSE", "Viscose", "Ring", None] + [None] * 10 + [0.02 if s == "W" else None for s in S4_STAGES]
    tr = [
        4,
        "TR (65/35)",
        "",
        "Ring",
        None,
        "PE Normal - Ring",
        "Viscose - ring",
        None,
        None,
        None,
        0.65,
        0.35,
        None,
        None,
        None,
    ]
    tr += [None] * len(S4_STAGES)
    for row in (pe, vi, tr):
        ws.append(row)
    if duplicate_group:
        w = 5 + 10 + S4_STAGES.index("W")  # the two LINEN rows differ in column W (not only in C3.x)
        ws.append([5, "LINEN", "Linen", "Ring", *([None] * (w - 4)), 0.01])
        ws.append([6, " linen ", "Linen", "Ring", *([None] * (w - 4)), 0.02])
    ws = sheet("1. Input NSM")
    ws.append(["Stt", "Mã sợi.", "Tên sợi", "Nhóm sợi"])
    ws.append(
        [
            "Stt",
            "Mã sợi",
            "Tên sợi",
            "Nhóm sợi",
            "Chi số",
            *SLOTS,
            "Type 3",
            *S4_STAGES,
            "RW",
            "Nhóm tiêu hao",
            "Nguyên liệu",
            "Loại đóng gói",
        ]
    )

    def nsm_row(code: str | None, name: str, blend: dict, nsm: dict, packaging: str = "Type 3") -> list:
        return [
            1,
            code,
            name,
            "x",
            "40",
            *(blend.get(s) for s in SLOTS),
            273,
            *(nsm.get(s, 0) for s in S4_STAGES),
            0,
            "g",
            "g",
            packaging,
        ]

    blend = {"Cotton": 0.7, "Other - C3.3": 0.3}
    nsm = {k: float(v) for k, v in S4_FILE_NSM.items()}
    ws.append(nsm_row("S4-2017", S4_NAME, blend, nsm))
    ws.append(nsm_row("S4 - 2017 (ĐCNSM2)", S4_NAME, blend, {**nsm, "V": 30}))
    ws.append(nsm_row(None, "CORE 30/1 (COTTON/ FDY 30D/24F/1) (70/30) MVS", blend, nsm))
    ws.append(nsm_row("S42003000010000000", "TR 30/1 (65/35)", {"PE": 0.65, "Other - C3.1": 0.35}, {"C2": 45, "W": 50}))
    ws.append(nsm_row("S4-2018", "BAD NAME", blend, nsm))
    for row in extra_nsm or []:
        ws.append(row)
    ws = sheet("4. Bang gia da tinh")
    ws.append([])
    ws.append([])
    ws.append(
        [
            "STT",
            "Mã Kế Toán",
            "Tên Sợi",
            "Chi phí",
            None,
            None,
            None,
            None,
            None,
            "Ngày tính giá",
            "SO check",
            "Mã Kế toán",
        ]
    )
    ws.append(
        [
            "STT",
            "Mã Kế Toán",
            "Tên Sợi",
            "Nhân công",
            "Khấu hao",
            "Điện",
            "Vật tư",
            "CPQL",
            "TC",
            "Ngày tính giá",
            "SO check",
            "Mã Kế toán",
        ]
    )
    ws.append(list(range(30, 42)))
    ws.append([1, "S4-2017", S4_NAME, 3618.95, 3428.98, 9835.53, 1554.20, 1340.17, 19777.84, "06/01/2024", "x", None])
    ws.append([2, "S42003000010000000", "TR 30/1 (65/35)", 1, 1, 1, 1, 1, 5, "06/01/2024", "x", None])
    ws.append([3, "S4-9999", "X", 1, 1, 1, 1, 1, 5, None, None, None])
    ws.append([4, "S4-2017 (ĐCNSM2)", S4_NAME, 1, 1, 1, 1, 1, 5, None, None, "S42KT"])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def run(content: bytes, commit: bool = True, mode: ImportMode = ImportMode.BAN_DAU, actor=SO, name: str = "s4.xlsx"):
    return ImportWorkbookService().handle(
        ImportWorkbookRequest(
            actor=actor,
            factory_code=S4,
            mode=mode,
            cutoff_date=date(2026, 10, 1),
            file=UploadedFileDTO(filename=name, content=content),
            commit=commit,
        )
    )


def test_initial_import_prices_and_compares() -> None:
    preview = run(s4_workbook(duplicate_group=True), commit=False)
    assert not preview.committed and preview.valid_count > 0
    reasons = {r.reason for r in preview.rejected_rows}
    assert "Nhóm 'LINEN' bị trùng trong file." in reasons  # AC-47
    assert any("chi số dạng" in r for r in reasons) and "Không có dữ liệu NSM cho mã 'S4-9999'." in reasons
    res = run(s4_workbook(duplicate_group=True))
    assert res.committed and res.import_id == "NF-S4-001" and res.result_file is not None
    assert "Đã tính giá" in res.message
    detail = GetYarnDetailService().handle(GetYarnDetailRequest(yarn_code="S4-2017"))
    assert abs(detail.yarn.total_vnd - D("19777.8356")) < D("0.0001")  # AC-01 through the import
    assert detail.yarn.first_priced_at == date(2024, 1, 6) and detail.yarn.effective_date == date(2026, 10, 1)
    revision = GetYarnDetailService().handle(GetYarnDetailRequest(yarn_code="S4-2017 (ĐCNSM2)"))  # AC-26
    assert revision.yarn.base_code == "S4-2017" and revision.yarn.revision_no == 2 and revision.yarn.is_locked
    assert (
        GetYarnDetailService().handle(GetYarnDetailRequest(yarn_code="S4-2020")).yarn.yarn_name.startswith("CORE 30/1")
    )
    kept = GetYarnDetailService().handle(GetYarnDetailRequest(yarn_code="S42003000010000000"))
    assert kept.yarn.is_locked and kept.yarn.status == YarnStatus.CO_MA_KE_TOAN
    assert kept.lines[0].accounting_code == "S42003000010000000"
    groups = {g.group_key: g for g in LossGroupRepository().for_factory("S4")}
    assert groups["TR (65/35)"].level == LossGroupLevel.THU_CAP and groups["TR (65/35)"].losses["W"] == D(
        "0.0135"
    )  # AC-20
    comparison = GetLegacyComparisonService().handle(GetLegacyComparisonRequest(import_id="NF-S4-001"))
    assert comparison.matched_count >= 1 and comparison.file.content[:2] == b"PK"  # AC-02
    mismatch = GetLegacyComparisonService().handle(
        GetLegacyComparisonRequest(import_id="NF-S4-001", only_mismatch=True)
    )
    assert all(not r.matched for r in mismatch.items)
    batches = ListImportBatchesService().handle(ListImportBatchesRequest(factory_code=S4)).items
    assert batches[0].import_id == "NF-S4-001" and batches[0].rejected_count >= 3


def test_file_type_sheets_and_permissions() -> None:
    with pytest.raises(BusinessRuleViolation, match=r"\.xlsx, \.xlsm hoặc \.xls"):  # AC-27
        run(b"%PDF", name="gia.pdf")
    with pytest.raises(AppError, match="File thiếu sheet: Hoi phe"):
        run(s4_workbook(sheets={"Dinh muc (phut)", "1. Input NSM"}))
    run(s4_workbook())
    xs4 = add_member("XS4", "XS40000", "An")
    with pytest.raises(BusinessRuleViolation, match="không được thực hiện thao tác này"):
        run(s4_workbook(), actor=xs4)
    res = run(s4_workbook(), actor=xs4, mode=ImportMode.HANG_LOAT, commit=False)
    skipped = {s.sheet: s.skipped_message for s in res.sheets}
    assert skipped["Dinh muc (phut)"] == "Sheet 'Dinh muc (phut)' chỉ Phòng SO được cập nhật, đã bỏ qua."
    assert skipped["1. Input NSM"] == ""
    locked = {r.reason for r in res.rejected_rows if r.sheet == "1. Input NSM"}
    assert any("đã có mã kế toán" in r for r in locked)  # BR-44: locked yarns are not updated


def test_template_round_trip() -> None:
    run(s4_workbook())
    file = DownloadTemplateService().handle(DownloadTemplateRequest(factory_code=S4))
    assert file.filename == "mau-du-lieu-tinh-gia-S4.xlsx"
    book = load_workbook(BytesIO(file.content))
    assert book.sheetnames == [
        "Dinh muc (phut)",
        "Hoi phe",
        "Dong goi",
        "QTCN theo nhom soi",
        "2. Input tieu hao",
        "1. Input NSM",
        "4. Bang gia da tinh",
    ]
    header = [c.value for c in book["1. Input NSM"][1]]
    assert header[:3] == ["Stt", "Mã sợi", "Tên sợi"] and "S Compact" in header and "Cotton" in header
    again = run(file.content, mode=ImportMode.HANG_LOAT, name="mau.xlsx")
    assert again.committed and sum(s.rejected for s in again.sheets if s.sheet != "1. Input NSM") == 0


def test_template_s1_columns() -> None:
    file = DownloadTemplateService().handle(DownloadTemplateRequest(factory_code=FactoryCode.S1))
    header = [c.value for c in load_workbook(BytesIO(file.content))["1. Input NSM"][1]]
    assert header[3:6] == ["PE", "Viscose/other", "Cotton"]  # AC-41
    assert header[6:17] == ["C1", "C2", "PD", "PD2", "MD1", "MD2", "MD3", "R", "S", "W", "V"]
