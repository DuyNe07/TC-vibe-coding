"""2026-09-28: default packaging (BR-60), multi-pass carding lines (BR-56..BR-59), import / template / old stored data."""

import json
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.core.config import get_settings
from backend.core.exceptions import BusinessRuleViolation
from backend.features.yarn_costing.dto.common_dto import YarnComponentDTO
from backend.features.yarn_costing.dto.import_dto import DownloadTemplateRequest
from backend.features.yarn_costing.dto.master_dto import PackagingRowDTO, SavePackagingTypesRequest
from backend.features.yarn_costing.dto.price_dto import ListPriceLinesRequest
from backend.features.yarn_costing.dto.pricing_dto import GetYarnDetailRequest, SaveYarnNsmRequest
from backend.features.yarn_costing.dto.reference_dto import GetReferenceDataRequest
from backend.features.yarn_costing.models.enums import FactoryCode, ImportMode, YarnStatus
from backend.features.yarn_costing.repositories.loss_group_repository import LossGroupRepository
from backend.features.yarn_costing.repositories.rate_version_repository import RateVersionRepository
from backend.features.yarn_costing.repositories.stage_repository import StageRepository
from backend.features.yarn_costing.repositories.yarn_repository import YarnRepository
from backend.features.yarn_costing.services.download_template_service import DownloadTemplateService
from backend.features.yarn_costing.services.get_reference_data_service import GetReferenceDataService
from backend.features.yarn_costing.services.get_yarn_detail_service import GetYarnDetailService
from backend.features.yarn_costing.services.list_price_lines_service import ListPriceLinesService
from backend.features.yarn_costing.services.save_packaging_types_service import SavePackagingTypesService
from backend.features.yarn_costing.services.save_yarn_nsm_service import SaveYarnNsmService
from backend.features.yarn_costing.tests.helpers import create_request, setup_factory
from backend.features.yarn_costing.tests.test_import_services import run, s4_workbook

D = Decimal
S4 = FactoryCode.S4


def test_default_packaging_is_exposed_and_single() -> None:
    setup_factory(S4)
    res = SavePackagingTypesService().handle(
        SavePackagingTypesRequest(
            factory_code=S4,
            rows=[
                PackagingRowDTO(code="Type 3", name="Xe", minutes_per_kg=D("0.2194564694564695"), is_default=True),
                PackagingRowDTO(code="Type 1", name="Bao", minutes_per_kg=D("0.3")),
            ],
        )
    )
    assert res.batch is None  # changing the default does not recalculate
    ref = GetReferenceDataService().handle(GetReferenceDataRequest(factory_code=S4))
    s4 = next(f for f in ref.factories if f.code == S4)
    assert s4.default_packaging_code == "Type 3" and s4.carding_stages == ["C1", "C2", "C3"]  # AC-50
    assert {p.code: p.is_default for p in ref.packaging} == {"Type 1": False, "Type 3": True}
    code = create_request("PE 30/1", packaging="Type 3")
    assert YarnRepository().find_one(lambda y: y.yarn_code == code).packaging_code == "Type 3"
    with pytest.raises(BusinessRuleViolation, match="Mỗi xưởng chỉ được chọn một loại đóng gói mặc định."):
        SavePackagingTypesService().handle(
            SavePackagingTypesRequest(
                factory_code=S4,
                rows=[
                    PackagingRowDTO(code="Type 3", name="Xe", minutes_per_kg=D("0.2194564694564695"), is_default=True),
                    PackagingRowDTO(code="Type 1", name="Bao", minutes_per_kg=D("0.3"), is_default=True),
                ],
            )
        )


def test_import_builds_components_and_keeps_the_price() -> None:
    run(s4_workbook())
    detail = GetYarnDetailService().handle(GetYarnDetailRequest(yarn_code="S4-2017"))
    assert abs(detail.yarn.total_vnd - D("19777.8356")) < D("0.0001")  # AC-55
    assert [(c.material_label, c.carding_stage, c.pass_no) for c in detail.components] == [
        ("Other - C3.3", "C3", 1),
        ("Cotton", "C1", 1),
    ]
    assert "C3" in {s.stage_code for s in detail.stage_details} and detail.carding_stages == ["C1", "C2", "C3"]
    assert all(s.stage_code not in {"C1", "C2", "C3"} for s in detail.stages)  # table ② = other stages only
    tr = GetYarnDetailService().handle(GetYarnDetailRequest(yarn_code="S42003000010000000"))
    # BR-58: secondary group → material_tech_key of the component in the same cell (PE → C2, Other - C3.1 → C3)
    assert [(c.material_label, c.carding_stage) for c in tr.components] == [
        ("PE Normal - Ring", "C2"),
        ("Viscose - Ring", "C3"),
    ]
    rows = ListPriceLinesService().handle(ListPriceLinesRequest(factory_code=S4)).items
    text = next(r.components_text for r in rows if r.line_code == "S4-2017")
    assert text == "Other - C3.3 30% (C3 lượt 1); Cotton 70% (C1)"


def test_import_merges_rows_that_differ_only_in_c3_columns() -> None:
    content = s4_workbook()
    book = load_workbook(BytesIO(content))
    ws = book["2. Input tieu hao"]
    header = [c.value for c in ws[4]]
    row_a = [None] * len(header)
    row_b = [None] * len(header)
    for row, col, value in ((row_a, "C3.1", 0.034), (row_b, "C3.2", 0.035)):
        row[header.index("Nhóm sợi")] = "MODAL"
        row[header.index("Nguyên liệu")] = "Modal"
        row[header.index("Công nghệ")] = "Ring"
        row[header.index("W")] = 0.01
        row[header.index(col)] = value
    ws.append(row_a)
    ws.append(row_b)
    buffer = BytesIO()
    book.save(buffer)
    res = run(buffer.getvalue())
    assert not any(r.key == "MODAL" for r in res.rejected_rows)
    modal = next(g for g in LossGroupRepository().for_factory("S4") if g.group_key == "MODAL")
    assert modal.losses["C3"] == D("0.034") and "C3.1" not in modal.losses  # AC-56


def test_template_keeps_the_file_layout_and_remark() -> None:
    run(s4_workbook())
    code = create_request("CORE 20/1 (COTTON/ FDY 30D/24F/1) (70/30) MVS")
    SaveYarnNsmService().handle(
        SaveYarnNsmRequest(
            yarn_code=code,
            nsm={"CM": D(50), "MD1": D(100), "MD2": D(90), "MD3": D(83), "V": D(25)},
            components=[
                YarnComponentDTO(material_label="Cotton", ratio=D("0.7"), carding_stage="C1", nsm_carding=D(45)),
                YarnComponentDTO(material_label="A", ratio=D("0.15"), carding_stage="C2", nsm_carding=D(45)),
                YarnComponentDTO(material_label="B", ratio=D("0.15"), carding_stage="C2", nsm_carding=D(45)),
            ],
            nsm_warning_reason="Pha thêm 2 nguyên liệu",
            save=True,
        )
    )
    file = DownloadTemplateService().handle(DownloadTemplateRequest(factory_code=S4))
    book = load_workbook(BytesIO(file.content))
    loss_header = [c.value for c in book["2. Input tieu hao"][1]]
    assert "C3" in loss_header and "C3.1" not in loss_header  # AC-59
    nsm_sheet = book["1. Input NSM"]
    nsm_header = [c.value for c in nsm_sheet[1]]
    assert {"C3.1", "C3.2", "C3.3", "REMARK"} <= set(nsm_header) and "C3" not in nsm_header
    rows = {r[1]: dict(zip(nsm_header, r, strict=False)) for r in nsm_sheet.iter_rows(min_row=2, values_only=True)}
    assert rows[code]["REMARK"] == "Có thành phần chỉ sửa được trên màn hình"
    assert rows["S4-2017"]["C3.1"] == 45 and rows["S4-2017"]["Other - C3.1"] == 0.3  # C3 lượt 1 → C3.1
    before = YarnRepository().find_one(lambda y: y.yarn_code == code).components
    run(file.content, mode=ImportMode.HANG_LOAT, name="mau.xlsx")
    after = YarnRepository().find_one(lambda y: y.yarn_code == code).components
    assert [c.material_label for c in after] == [c.material_label for c in before]  # BR-58: REMARK keeps them
    kept = YarnRepository().find_one(lambda y: y.yarn_code == "S4-2017").components
    assert [c.material_label for c in kept] == ["Other - C3.3", "Cotton"]  # names survive the round trip


def test_old_stored_data_still_loads() -> None:
    folder = get_settings().feature_data_dir("yarn_costing")
    folder.mkdir(parents=True, exist_ok=True)
    stage = {
        "factory_code": "S4",
        "stage_group": "C3.1",
        "recovery_type": "Chải thô",
        "is_merge_stage": False,
        "secondary_method": "Cộng thẳng",
        "varies_with_count": False,
        "rate_code_source": "Chính công đoạn",
        "is_packaging": False,
    }
    stages = [
        {
            **stage,
            "id": "a",
            "stage_code": "C1",
            "stage_name": "C1",
            "stage_group": "C1",
            "sequence": 1,
            "blend_slot": "Cotton",
            "cumulative_stages": ["SBD"],
        },
        {
            **stage,
            "id": "b",
            "stage_code": "C3.1",
            "stage_name": "x",
            "sequence": 3,
            "blend_slot": "Other - C3.1",
            "cumulative_stages": [],
        },
        {
            **stage,
            "id": "c",
            "stage_code": "C3.2",
            "stage_name": "x",
            "sequence": 4,
            "blend_slot": "Other - C3.2",
            "cumulative_stages": [],
        },
        {
            **stage,
            "id": "d",
            "stage_code": "SBD",
            "stage_name": "SBD",
            "stage_group": "SBD",
            "sequence": 6,
            "blend_slot": "Cotton",
            "is_merge_stage": True,
            "cumulative_stages": [],
        },
        {
            **stage,
            "id": "e",
            "stage_code": "F",
            "stage_name": "F",
            "stage_group": "F",
            "sequence": 21,
            "blend_slot": "Không",
            "is_packaging": True,
            "recovery_type": "Không",
            "cumulative_stages": [],
        },
    ]
    (folder / "stages.json").write_text(json.dumps(stages), encoding="utf-8")
    (folder / "loss_groups.json").write_text(
        json.dumps(
            [
                {
                    "id": "g",
                    "factory_code": "S4",
                    "group_key": "PE 30/1",
                    "yarn_group_label": "PE",
                    "technology": "Ring",
                    "losses": {"C3.1": "0.035", "C3.2": "0.034", "C1": "0.07"},
                }
            ]
        ),
        encoding="utf-8",
    )
    line = {"labor": "1", "depreciation": "1", "electricity": "1", "materials": "1", "overhead": "1"}
    (folder / "rate_versions.json").write_text(
        json.dumps(
            [
                {
                    "id": "v",
                    "version_code": "DM-S4-001",
                    "factory_code": "S4",
                    "effective_date": "2026-01-01",
                    "reason": "x",
                    "lines": [{**line, "rate_code": "C3.1"}, {**line, "rate_code": "C1"}],
                }
            ]
        ),
        encoding="utf-8",
    )
    yarn = {
        "id": "y",
        "factory_code": "S4",
        "yarn_code": "S4-1",
        "base_code": "S4-1",
        "yarn_name": "PE 30/1",
        "loss_group_key": "PE",
        "status": "Chưa tính được",
        "blend_ratios": {"Cotton": "0.7", "Other - C3.2": "0.3"},
        "nsm": {"C1": "45", "C3.2": "40", "V": "25"},
        "stage_details": [{"sequence": 4, "stage_code": "C3.2", "stage_name": "x", "run_minutes": "1"}],
    }
    (folder / "yarns.json").write_text(json.dumps([yarn]), encoding="utf-8")
    (folder / "packaging_types.json").write_text(
        json.dumps([{"id": "p", "factory_code": "S4", "code": "Type 3", "name": "x", "minutes_per_kg": "0.2"}]),
        encoding="utf-8",
    )
    codes = [(s.stage_code, s.sequence, s.blend_slot, s.is_carding) for s in StageRepository().for_factory("S4")]
    assert codes == [
        ("C1", 1, "Tỷ lệ của lượt", True),
        ("C3", 2, "Tỷ lệ của lượt", True),
        ("SBD", 3, "Tỷ lệ dây chuyền C1", False),
        ("F", 4, "Không", False),
    ]
    assert LossGroupRepository().list_all()[0].losses == {"C3": D("0.034"), "C1": D("0.07")}
    assert RateVersionRepository().current("S4").line_for("C3") is not None
    loaded = YarnRepository().list_all()[0]
    assert [(c.material_label, c.carding_stage, c.nsm_carding) for c in loaded.components] == [
        ("Other - C3.2", "C3", D(40)),
        ("Cotton", "C1", D(45)),
    ]
    assert loaded.nsm == {"V": D(25)} and loaded.status == YarnStatus.CHUA_TINH_DUOC
    assert (loaded.stage_details[0].stage_code, loaded.stage_details[0].pass_no) == ("C3", 2)
    ref = GetReferenceDataService().handle(GetReferenceDataRequest(factory_code=S4))
    assert ref.packaging[0].is_default is False
